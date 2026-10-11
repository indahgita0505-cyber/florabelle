import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_florabelle_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Pastel & Cute Aesthetic
    PINK_MAIN = RGBColor(244, 114, 182)       # #f472b6 (Rose Pink)
    PINK_DARK = RGBColor(219, 39, 119)       # #db2777 (Deep Rose)
    PINK_BG = RGBColor(255, 241, 246)         # #fff1f6 (Soft Blush Cream)
    WHITE = RGBColor(255, 255, 255)           # #ffffff
    PURPLE_SOFT = RGBColor(192, 132, 252)     # #c084fc (Lilac)
    PURPLE_DARK = RGBColor(147, 51, 234)      # #9333ea
    GREEN_SOFT = RGBColor(134, 239, 172)      # #86efac (Sage Green)
    GREEN_DARK = RGBColor(22, 163, 74)        # #16a34a
    AMBER_SOFT = RGBColor(253, 230, 138)      # #fde68a (Honey Yellow)
    AMBER_DARK = RGBColor(217, 119, 6)        # #d97706
    BLUE_SOFT = RGBColor(186, 230, 253)       # #bae6fd (Sky Blue)
    BLUE_DARK = RGBColor(2, 132, 199)         # #0284c7
    TEXT_DARK = RGBColor(51, 65, 85)          # #334155 (Slate 700)
    TEXT_MUTED = RGBColor(100, 116, 139)      # #64748b (Slate 500)
    BORDER_PINK = RGBColor(251, 207, 232)     # #fbcfe8

    def set_slide_background(slide, color=PINK_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background() # no line
        return bg

    def add_header(slide, title_text, category_tag="🌸 FLORABELLE.ID PALU"):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = PINK_DARK
        p_tag.font.name = "Segoe UI"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.font.name = "Georgia"

        # Decorative line
        line = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(1.5), Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = PINK_MAIN
        line.line.fill.background()

    blank_layout = prs.slide_layouts[6]

    # ══════════════════════════════════════════════════════════════
    # SLIDE 1: COVER
    # ══════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, RGBColor(255, 245, 248))

    # Center Aesthetic Card
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = WHITE
    card1.line.color.rgb = BORDER_PINK
    card1.line.width = Pt(2)

    # Top Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.666), Inches(1.5), Inches(4.0), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(253, 242, 248)
    badge.line.color.rgb = PINK_MAIN
    badge.line.width = Pt(1.5)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "✨ PROPOSAL & DOKUMENTASI SISTEM ✨"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = PINK_DARK
    p_b.font.name = "Segoe UI"
    p_b.alignment = PP_ALIGN.CENTER

    # Title
    t_box = s1.shapes.add_textbox(Inches(2.0), Inches(2.2), Inches(9.333), Inches(1.4))
    tf_t = t_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "🌸 Florabelle.id Palu"
    p_t.font.size = Pt(40)
    p_t.font.bold = True
    p_t.font.color.rgb = PINK_DARK
    p_t.font.name = "Georgia"
    p_t.alignment = PP_ALIGN.CENTER

    # Subtitle
    sub_box = s1.shapes.add_textbox(Inches(2.0), Inches(3.5), Inches(9.333), Inches(0.8))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Platform E-Commerce Florist & Gift Shop Modern Berbasis Web\ndi Kota Palu, Sulawesi Tengah"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = TEXT_DARK
    p_sub.font.name = "Segoe UI"
    p_sub.alignment = PP_ALIGN.CENTER

    # Features Pill Row
    pills = [
        ("🌹 Fresh & Dried Flowers", PINK_BG, PINK_DARK),
        ("💌 Custom Greeting Card", RGBColor(243, 232, 255), PURPLE_DARK),
        ("🛵 Grab Express Palu", RGBColor(240, 253, 244), GREEN_DARK),
    ]
    pill_w = Inches(2.8)
    start_x = Inches(2.3)
    for i, (txt, bg_c, fg_c) in enumerate(pills):
        p_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x + (i * Inches(3.0)), Inches(4.5), pill_w, Inches(0.45))
        p_shape.fill.solid()
        p_shape.fill.fore_color.rgb = bg_c
        p_shape.line.fill.background()
        tf_p = p_shape.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = txt
        p_p.font.size = Pt(11)
        p_p.font.bold = True
        p_p.font.color.rgb = fg_c
        p_p.font.name = "Segoe UI"
        p_p.alignment = PP_ALIGN.CENTER

    # Footer note
    f_box = s1.shapes.add_textbox(Inches(2.0), Inches(5.6), Inches(9.333), Inches(0.5))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Crafted with Love & Technology 💖 · Laravel 11/13 + React Inertia + Tailwind CSS"
    p_f.font.size = Pt(11)
    p_f.font.color.rgb = TEXT_MUTED
    p_f.font.name = "Segoe UI"
    p_f.alignment = PP_ALIGN.CENTER


    # ══════════════════════════════════════════════════════════════
    # SLIDE 2: LATAR BELAKANG
    # ══════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "🌿 Latar Belakang Dibuatnya Florabelle.id", "Konteks & Masalah")

    cards_data_s2 = [
        ("📍 Kebutuhan Florist di Kota Palu", 
         "Tingginya antusiasme masyarakat Palu untuk merayakan momen spesial (Wisuda Untad/kampus lokal, Ulang Tahun, Anniversary, Sempro, & Event Bisnis) yang membutuhkan hadiah buket bunga berkualitas.",
         PINK_BG, PINK_DARK, BORDER_PINK),
        
        ("📱 Kendala Pemesanan Manual (DM/WA)", 
         "Pemesanan tradisional via chat rawan terjadi salah rekap alamat, keterlambatan respons admin saat jam sibuk, harga ongkir yang tidak transparan, dan kartu ucapan yang sering tertukar.",
         RGBColor(254, 242, 242), RGBColor(220, 38, 38), RGBColor(254, 202, 202)),
        
        ("🛍️ Solusi Digital Florabelle.id", 
         "Menghadirkan website e-commerce modern 24/7 dengan katalog interaktif, kustomisasi buket, catatan kartu ucapan otomatis, perhitungan ongkir lokal akurat, dan pelacakan pesanan.",
         RGBColor(240, 253, 244), GREEN_DARK, RGBColor(187, 247, 208))
    ]

    for i, (title, desc, bg_c, text_c, border_c) in enumerate(cards_data_s2):
        x = Inches(0.8 + (i * 3.95))
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.75), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_c
        card.line.color.rgb = border_c
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.4)
        tf.margin_left = Inches(0.35)
        tf.margin_right = Inches(0.35)

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = text_c
        p1.font.name = "Georgia"
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK
        p2.font.name = "Segoe UI"
        p2.line_spacing = 1.3


    # ══════════════════════════════════════════════════════════════
    # SLIDE 3: TUJUAN PENGEMBANGAN
    # ══════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "🎯 Tujuan Pengembangan Website", "Visi & Target Sistem")

    goals = [
        ("01", "Digitalisasi & Transaksi 24/7", 
         "Memberikan akses belanja buket bunga kapan saja tanpa harus menunggu balasan chat manual admin toko.",
         PINK_BG, PINK_DARK),
        ("02", "Personalisasi Pesan & Buket", 
         "Menyediakan fitur gratis kartu ucapan (Greeting Card) & request custom kertas/pita buket sesuai selera.",
         RGBColor(243, 232, 255), PURPLE_DARK),
        ("03", "Otomatisasi Stok & Keranjang", 
         "Sinkronisasi database realtime untuk menghindari overselling (stok habis) dan kemudahan kalkulasi total.",
         RGBColor(254, 243, 199), AMBER_DARK),
        ("04", "Integrasi Pengiriman Kota Palu", 
         "Menghitung estimasi tarif Grab Express Palu secara akurat untuk pengiriman kilat tepat waktu ke customer.",
         RGBColor(240, 253, 244), GREEN_DARK),
    ]

    for i, (num, title, desc, bg_c, text_c) in enumerate(goals):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + (col * 5.95))
        y = Inches(2.0 + (row * 2.45))

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.25))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_PINK
        card.line.width = Pt(1.5)

        # Number pill inside
        np = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.3), y + Inches(0.3), Inches(0.75), Inches(0.75))
        np.fill.solid()
        np.fill.fore_color.rgb = bg_c
        np.line.fill.background()
        tf_n = np.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(16)
        p_n.font.bold = True
        p_n.font.color.rgb = text_c
        p_n.alignment = PP_ALIGN.CENTER

        # Text box next to number
        tb = s3.shapes.add_textbox(x + Inches(1.2), y + Inches(0.2), Inches(4.2), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p1.font.name = "Georgia"
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Segoe UI"
        p2.line_spacing = 1.2


    # ══════════════════════════════════════════════════════════════
    # SLIDE 4: TECH STACK
    # ══════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "💻 Tech Stack & Arsitektur Website", "Teknologi & Infrastruktur")

    tech_data = [
        ("🐘 Laravel 11/13", "Backend Framework", 
         "Menyediakan arsitektur MVC yang kokoh, ORM Eloquent yang kuat, autentikasi aman, dan manajemen database relasional yang andal.",
         RGBColor(254, 226, 226), RGBColor(220, 38, 38)),
        
        ("⚡ Inertia.js + React 18", "Frontend Monolith SPA", 
         "Menghadirkan pengalaman Single Page Application yang super cepat dan responsif tanpa repot membangun REST API terpisah.",
         RGBColor(238, 242, 255), RGBColor(79, 70, 229)),
        
        ("🎨 Tailwind CSS 3/4", "Modern Styling & UI", 
         "Menciptakan antarmuka yang sangat cantik, mobile-first, bertema pastel floral dengan micro-interaction yang ramah pengguna.",
         RGBColor(224, 242, 254), RGBColor(2, 132, 199)),
        
        ("🗄️ MySQL & Vite", "Database & Tooling", 
         "Penyimpanan data produk, keranjang, & order terstruktur dengan Hot Module Replacement (HMR) berkecepatan tinggi saat development.",
         RGBColor(254, 243, 199), AMBER_DARK)
    ]

    for i, (title, sub, desc, bg_c, text_c) in enumerate(tech_data):
        x = Inches(0.8 + (i * 2.95))
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_PINK
        card.line.width = Pt(1.5)

        # Header box in card
        hb = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), Inches(2.2), Inches(2.4), Inches(1.1))
        hb.fill.solid()
        hb.fill.fore_color.rgb = bg_c
        hb.line.fill.background()
        tf_h = hb.text_frame
        tf_h.word_wrap = True
        p_h1 = tf_h.paragraphs[0]
        p_h1.text = title
        p_h1.font.size = Pt(14)
        p_h1.font.bold = True
        p_h1.font.color.rgb = text_c
        p_h1.font.name = "Georgia"
        p_h1.alignment = PP_ALIGN.CENTER
        
        p_h2 = tf_h.add_paragraph()
        p_h2.text = sub
        p_h2.font.size = Pt(10)
        p_h2.font.color.rgb = TEXT_DARK
        p_h2.alignment = PP_ALIGN.CENTER

        # Body desc
        tb = s4.shapes.add_textbox(x + Inches(0.2), Inches(3.4), Inches(2.4), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p_d = tf.paragraphs[0]
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_DARK
        p_d.font.name = "Segoe UI"
        p_d.line_spacing = 1.3


    # ══════════════════════════════════════════════════════════════
    # SLIDE 5: MANFAAT WEBSITE
    # ══════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "🎁 Manfaat Website untuk Semua Pihak", "Dampak & Value Proposition")

    benefits = [
        ("💖 Bagi Pelanggan (Customer)", [
            "Belanja praktis 24 jam tanpa perlu menunggu admin membalas chat.",
            "Katalog visual lengkap dengan foto berkualitas tinggi & filter kategori.",
            "Gratis kartu ucapan dengan input pesan langsung di keranjang belanja.",
            "Perhitungan harga dan ongkir Grab Palu yang transparan."
        ], PINK_BG, PINK_DARK),

        ("🏪 Bagi Pemilik Toko (Florist)", [
            "Pencatatan pesanan otomatis dan tersusun rapi di database.",
            "Mengurangi risiko human-error (salah catat nama, alamat, atau pesan kartu).",
            "Manajemen stok realtime sehingga terhindar dari stok kosong.",
            "Meningkatkan branding & kredibilitas toko bunga terdepan di Palu."
        ], RGBColor(243, 232, 255), PURPLE_DARK),

        ("🛵 Bagi Kurir & Logistik", [
            "Format alamat & nomor WhatsApp penerima tersimpan sangat jelas.",
            "Kemudahan koordinasi pengiriman instan di seluruh area Kota Palu.",
            "Efisiensi waktu pengantaran untuk menjaga kesegaran bunga."
        ], RGBColor(240, 253, 244), GREEN_DARK)
    ]

    for i, (title, bullets, bg_c, text_c) in enumerate(benefits):
        x = Inches(0.8 + (i * 3.95))
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.75), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_PINK
        card.line.width = Pt(1.5)

        # Header tag
        ht = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.25), Inches(2.2), Inches(3.25), Inches(0.6))
        ht.fill.solid()
        ht.fill.fore_color.rgb = bg_c
        ht.line.fill.background()
        tf_ht = ht.text_frame
        p_ht = tf_ht.paragraphs[0]
        p_ht.text = title
        p_ht.font.size = Pt(13)
        p_ht.font.bold = True
        p_ht.font.color.rgb = text_c
        p_ht.font.name = "Georgia"
        p_ht.alignment = PP_ALIGN.CENTER

        # Bullets
        tb = s5.shapes.add_textbox(x + Inches(0.25), Inches(2.9), Inches(3.25), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True
        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = "• " + b
            p.font.size = Pt(11.5)
            p.font.color.rgb = TEXT_DARK
            p.font.name = "Segoe UI"
            p.space_after = Pt(8)
            p.line_spacing = 1.2


    # ══════════════════════════════════════════════════════════════
    # SLIDE 6: SWOT (STRENGTHS & WEAKNESSES)
    # ══════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "📊 Analisis SWOT (Kekuatan & Kelemahan)", "Strategic Evaluation · Part 1")

    # Strengths Card
    c_str = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8))
    c_str.fill.solid()
    c_str.fill.fore_color.rgb = RGBColor(240, 253, 244)
    c_str.line.color.rgb = RGBColor(187, 247, 208)
    c_str.line.width = Pt(1.5)

    tf_str = c_str.text_frame
    tf_str.word_wrap = True
    tf_str.margin_top = Inches(0.35)
    tf_str.margin_left = Inches(0.4)
    tf_str.margin_right = Inches(0.4)

    p = tf_str.paragraphs[0]
    p.text = "💪 STRENGTHS (Kekuatan)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GREEN_DARK
    p.font.name = "Georgia"
    p.space_after = Pt(12)

    str_points = [
        ("🎨 Desain Modern & Estetik", "Tampilan pastel floral yang memikat generasi muda (Gen Z & Milenial)."),
        ("⚡ Sinkronisasi Dinamis", "State keranjang dan badge notifikasi realtime berbasis Laravel + Inertia."),
        ("💐 Variasi Produk Lengkap", "Menyediakan buket bunga segar, bunga kering (dried), bloom box, hingga bunga artifisial."),
        ("🛵 Spesialisasi Lokal Palu", "Optimalisasi pengiriman instan menjangkau seluruh kecamatan di Palu.")
    ]
    for title, desc in str_points:
        p1 = tf_str.add_paragraph()
        p1.text = "✓ " + title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p1.font.name = "Segoe UI"
        
        p2 = tf_str.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Segoe UI"
        p2.space_after = Pt(6)

    # Weaknesses Card
    c_wk = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8))
    c_wk.fill.solid()
    c_wk.fill.fore_color.rgb = RGBColor(254, 242, 242)
    c_wk.line.color.rgb = RGBColor(254, 202, 202)
    c_wk.line.width = Pt(1.5)

    tf_wk = c_wk.text_frame
    tf_wk.word_wrap = True
    tf_wk.margin_top = Inches(0.35)
    tf_wk.margin_left = Inches(0.4)
    tf_wk.margin_right = Inches(0.4)

    p = tf_wk.paragraphs[0]
    p.text = "⚠️ WEAKNESSES (Kelemahan)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(220, 38, 38)
    p.font.name = "Georgia"
    p.space_after = Pt(12)

    wk_points = [
        ("🌐 Ketergantungan Internet", "Pelanggan harus memiliki koneksi internet untuk mengakses katalog web."),
        ("🥀 Sifat Bunga Perishable", "Bunga segar memiliki masa simpan terbatas, menuntut akurasi perkiraan stok."),
        ("💳 Gateway Otomatis Masih Terbatas", "Saat ini pemesanan diarahkan ke WhatsApp/transfer manual sebelum integrasi full PG."),
        ("📱 Belum Ada Aplikasi Native", "Saat ini masih berbasis Progressive Web App (PWA)/Mobile Browser.")
    ]
    for title, desc in wk_points:
        p1 = tf_wk.add_paragraph()
        p1.text = "• " + title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p1.font.name = "Segoe UI"
        
        p2 = tf_wk.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Segoe UI"
        p2.space_after = Pt(6)


    # ══════════════════════════════════════════════════════════════
    # SLIDE 7: SWOT (OPPORTUNITIES & THREATS)
    # ══════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "📊 Analisis SWOT (Peluang & Ancaman)", "Strategic Evaluation · Part 2")

    # Opportunities Card
    c_opp = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8))
    c_opp.fill.solid()
    c_opp.fill.fore_color.rgb = RGBColor(238, 242, 255)
    c_opp.line.color.rgb = RGBColor(199, 210, 254)
    c_opp.line.width = Pt(1.5)

    tf_opp = c_opp.text_frame
    tf_opp.word_wrap = True
    tf_opp.margin_top = Inches(0.35)
    tf_opp.margin_left = Inches(0.4)
    tf_opp.margin_right = Inches(0.4)

    p = tf_opp.paragraphs[0]
    p.text = "🌟 OPPORTUNITIES (Peluang)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(79, 70, 229)
    p.font.name = "Georgia"
    p.space_after = Pt(12)

    opp_points = [
        ("🎓 Pasar Wisuda & Kampus Luas", "Ribuan mahasiswa di Untad, UIN Datokarama, dll. yang merayakan wisuda & yudisium berkala."),
        ("🤝 Kemitraan EO & Wedding Planner", "Peluang kerja sama dekorasi bunga panggung, standing flower, dan buket pengantin."),
        ("🎁 Ekspansi Kategori Produk", "Penambahan bundling cokelat, boneka wisuda, hampers hari raya (Idul Fitri/Natal), dan kue."),
        ("📈 Tren Belanja Online Palu", "Pertumbuhan adopsi e-commerce lokal di Sulawesi Tengah yang semakin pesat.")
    ]
    for title, desc in opp_points:
        p1 = tf_opp.add_paragraph()
        p1.text = "★ " + title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p1.font.name = "Segoe UI"
        
        p2 = tf_opp.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Segoe UI"
        p2.space_after = Pt(6)

    # Threats Card
    c_thr = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8))
    c_thr.fill.solid()
    c_thr.fill.fore_color.rgb = RGBColor(254, 243, 199)
    c_thr.line.color.rgb = RGBColor(253, 230, 138)
    c_thr.line.width = Pt(1.5)

    tf_thr = c_thr.text_frame
    tf_thr.word_wrap = True
    tf_thr.margin_top = Inches(0.35)
    tf_thr.margin_left = Inches(0.4)
    tf_thr.margin_right = Inches(0.4)

    p = tf_thr.paragraphs[0]
    p.text = "🛡️ THREATS (Ancaman)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER_DARK
    p.font.name = "Georgia"
    p.space_after = Pt(12)

    thr_points = [
        ("✈️ Pasokan & Logistik Bunga Impor", "Keterlambatan kargo penerbangan bunga dari Malang/Bandung dapat mempengaruhi stok."),
        ("🏪 Persaingan Florist Tradisional", "Kompetitor lama yang menawarkan harga perang murah tanpa sistem."),
        ("🌦️ Kondisi Cuaca Ekstrem", "Musim hujan di Kota Palu yang dapat menghambat kurir pengantaran instan."),
        ("📉 Fluktuasi Harga Bunga Saat Musim Puncak", "Kenaikan harga bahan baku bunga saat Valentine, Hari Ibu, dan musim wisuda.")
    ]
    for title, desc in thr_points:
        p1 = tf_thr.add_paragraph()
        p1.text = "▲ " + title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p1.font.name = "Segoe UI"
        
        p2 = tf_thr.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Segoe UI"
        p2.space_after = Pt(6)


    # ══════════════════════════════════════════════════════════════
    # SLIDE 8: KESIMPULAN & ROADMAP
    # ══════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "🎀 Kesimpulan & Rencana Pengembangan", "Penutup & Langkah Depan")

    conclusions = [
        ("🌸 Transformasi Digital Florist", 
         "Florabelle.id sukses memodernisasi cara warga Kota Palu memesan bunga dan kado spesial dengan antarmuka yang cepat, estetik, dan menyenangkan.",
         PINK_BG, PINK_DARK),
        
        ("💎 Efisiensi & Kepuasan Pelanggan", 
         "Mengatasi kendala salah rekap, stok realtime, dan pesan kartu ucapan yang rapi, menghasilkan pengalaman belanja bintang lima.",
         RGBColor(243, 232, 255), PURPLE_DARK),
        
        ("🚀 Roadmap Masa Depan", 
         "Integrasi Payment Gateway Otomatis (QRIS / Midtrans), pelacakan posisi kurir realtime, dan program membership poin loyalitas.",
         RGBColor(240, 253, 244), GREEN_DARK)
    ]

    for i, (title, desc, bg_c, text_c) in enumerate(conclusions):
        x = Inches(0.8 + (i * 3.95))
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.75), Inches(3.4))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_c
        card.line.color.rgb = BORDER_PINK
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.3)
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = text_c
        p1.font.name = "Georgia"
        p1.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK
        p2.font.name = "Segoe UI"
        p2.line_spacing = 1.3

    # Big Banner at bottom
    bb = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.3))
    bb.fill.solid()
    bb.fill.fore_color.rgb = WHITE
    bb.line.color.rgb = PINK_MAIN
    bb.line.width = Pt(2)

    tf_bb = bb.text_frame
    tf_bb.word_wrap = True
    p_bb1 = tf_bb.paragraphs[0]
    p_bb1.text = "“Florabelle.id — Mekarkan Setiap Momen Berharga di Kota Palu 💖”"
    p_bb1.font.size = Pt(18)
    p_bb1.font.bold = True
    p_bb1.font.color.rgb = PINK_DARK
    p_bb1.font.name = "Georgia"
    p_bb1.alignment = PP_ALIGN.CENTER
    p_bb1.space_after = Pt(4)

    p_bb2 = tf_bb.add_paragraph()
    p_bb2.text = "Terima Kasih! ✨ Let's Bloom Together with Florabelle.id ✨"
    p_bb2.font.size = Pt(12)
    p_bb2.font.color.rgb = TEXT_MUTED
    p_bb2.font.name = "Segoe UI"
    p_bb2.alignment = PP_ALIGN.CENTER

    output_path = os.path.abspath("Florabelle_Palu_Presentation.pptx")
    prs.save(output_path)
    print(f"SUCCESS: Presentation saved at {output_path}")

if __name__ == "__main__":
    create_florabelle_presentation()
